# /// script
# requires-python = ">=3.12"
# dependencies = ["python-docx==1.2.0"]
# ///
"""Spike 0.1 — generate test documents to check that tokens survive translation.

Usage:
    uv run spikes/token_survival.py [--out DIR]

Writes ``token_survival.docx``, ``token_survival.md`` and
``token_survival.expected.json`` (expected occurrence count per token) into DIR
(default: ``spikes/out``). Run the documents through the translation tool, then
check the result with ``spikes/token_survival_check.py``.
"""

from __future__ import annotations

import argparse
import json
from collections import Counter
from pathlib import Path

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

P = "⟦P-K7M2X⟧"
O = "⟦O-R4N8Q⟧"
E = "⟦E-W3X9Z⟧"
TOKENS = (P, O, E)

# 20 sentences, tokens at the start, in the middle and at the end.
SENTENCES: list[tuple[str, str]] = [
    ("fr", f"{P} a présenté le budget lors de la réunion de lundi."),
    ("fr", f"Le contrat a été signé par {P} au nom de {O} la semaine dernière."),
    ("fr", f"Pour toute question, merci d'écrire à {E}."),
    ("fr", f"{O} prévoit d'ouvrir un nouveau bureau l'année prochaine."),
    (
        "fr",
        f"Selon {P}, le projet sera livré en retard ; contactez {E} pour le détail.",
    ),
    ("fr", f"La facture doit être envoyée à {O}."),
    (
        "fr",
        f"Nous avons rencontré {P}, qui nous a recommandé de contacter {P} de nouveau en mars.",
    ),
    ("de", f"{P} hat den Bericht gestern an die Geschäftsführung geschickt."),
    ("de", f"Die Firma {O} hat ihren Sitz in einer kleinen Stadt."),
    ("de", f"Bitte senden Sie alle Unterlagen an {E}."),
    ("de", f"{O} und {P} haben die Vereinbarung gemeinsam unterzeichnet."),
    ("de", f"Laut {P} ist die Lieferung bereits unterwegs."),
    ("de", f"Die Anfrage wurde von {E} an {O} weitergeleitet."),
    ("de", f"Wir danken {P}."),
    ("en", f"{P} will chair the steering committee next quarter."),
    ("en", f"The proposal from {O} was rejected by the board."),
    ("en", f"Questions about the invoice should be sent to {E}."),
    ("en", f"{E} is monitored during office hours only."),
    ("en", f"After meeting {P} and the team at {O}, we agreed on a new schedule."),
    ("en", f"The final decision rests with {P}."),
]

TABLE_ROWS: list[tuple[str, str, str]] = [
    ("Nom / Name", "Organisation", "E-mail"),
    (P, O, E),
    (f"Contact : {P}", f"Partenaire {O}", f"Écrire à {E}"),
]

FOOTNOTE_TEXT = f"Source : entretien avec {P} ({O}), joignable à {E}."

FOOTNOTES_CT = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"
)


def expected_counts() -> dict[str, int]:
    text = "\n".join(s for _, s in SENTENCES)
    text += "\n" + "\n".join(" ".join(r) for r in TABLE_ROWS)
    text += "\n" + FOOTNOTE_TEXT
    counts = Counter({t: text.count(t) for t in TOKENS})
    return dict(counts)


def _add_footnote(doc: Document, paragraph_index: int, text: str) -> None:
    """python-docx has no footnote API: build word/footnotes.xml by hand."""
    footnotes_xml = (
        f"<w:footnotes {nsdecls('w')}>"
        '<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
        '<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:r><w:continuationSeparator/></w:r></w:p></w:footnote>'
        '<w:footnote w:id="1"><w:p>'
        '<w:r><w:rPr><w:vertAlign w:val="superscript"/></w:rPr><w:footnoteRef/></w:r>'
        f'<w:r><w:t xml:space="preserve"> {text}</w:t></w:r>'
        "</w:p></w:footnote>"
        "</w:footnotes>"
    )
    part = Part(
        PackURI("/word/footnotes.xml"),
        FOOTNOTES_CT,
        footnotes_xml.encode("utf-8"),
        doc.part.package,
    )
    doc.part.relate_to(part, RT.FOOTNOTES)
    ref_run = parse_xml(
        f'<w:r {nsdecls("w")}><w:rPr><w:vertAlign w:val="superscript"/></w:rPr>'
        '<w:footnoteReference w:id="1"/></w:r>'
    )
    doc.paragraphs[paragraph_index]._p.append(ref_run)


def write_docx(path: Path) -> None:
    doc = Document()
    doc.add_heading("Token survival test", level=1)
    for _, sentence in SENTENCES:
        doc.add_paragraph(sentence)
    table = doc.add_table(rows=len(TABLE_ROWS), cols=3)
    table.style = "Table Grid"
    for r, row in enumerate(TABLE_ROWS):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    # Footnote attached to the first body sentence (index 1: index 0 is the heading).
    _add_footnote(doc, 1, FOOTNOTE_TEXT)
    doc.save(str(path))


def write_md(path: Path) -> None:
    lines = ["# Token survival test", ""]
    for i, (_, sentence) in enumerate(SENTENCES):
        lines.append(sentence + ("[^1]" if i == 0 else ""))
        lines.append("")
    header, *rows = TABLE_ROWS
    lines.append("| " + " | ".join(header) + " |")
    lines.append("|---|---|---|")
    for row in rows:
        lines.append("| " + " | ".join(row) + " |")
    lines.append("")
    lines.append(f"[^1]: {FOOTNOTE_TEXT}")
    lines.append("")
    path.write_text("\n".join(lines), encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--out", type=Path, default=Path(__file__).parent / "out")
    args = parser.parse_args()
    out: Path = args.out
    out.mkdir(parents=True, exist_ok=True)
    write_docx(out / "token_survival.docx")
    write_md(out / "token_survival.md")
    (out / "token_survival.expected.json").write_text(
        json.dumps(expected_counts(), ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    print(
        f"Written to {out}: token_survival.docx, token_survival.md, token_survival.expected.json"
    )
    print("Expected occurrences:", expected_counts())


if __name__ == "__main__":
    main()
