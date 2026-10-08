# /// script
# dependencies = ["python-docx", "python-pptx", "openpyxl"]
# ///
# Validation structurelle des fichiers réécrits (ouverture par des bibliothèques tierces + XML bien formé).
import sys, zipfile
from xml.dom import minidom
import docx, pptx, openpyxl

ok = True
for path in sys.argv[1:]:
    try:
        with zipfile.ZipFile(path) as z:
            assert z.testzip() is None
            for n in z.namelist():
                if n.endswith((".xml", ".rels")):
                    minidom.parseString(z.read(n))
        if path.endswith(".docx"):
            d = docx.Document(path); txt = "\n".join(p.text for p in d.paragraphs)
        elif path.endswith(".pptx"):
            pr = pptx.Presentation(path); txt = "\n".join(sh.text_frame.text for s in pr.slides for sh in s.shapes if sh.has_text_frame)
        else:
            wb = openpyxl.load_workbook(path); txt = "\n".join(str(c.value) for ws in wb for r in ws.iter_rows() for c in r if c.value is not None)
        print(f"OK   {path} : token présent = {'⟦P-K7M2X⟧' in txt}")
    except Exception as e:
        ok = False
        print(f"KO   {path} : {e!r}")
sys.exit(0 if ok else 1)
