# /// script
# dependencies = ["python-docx", "python-pptx", "openpyxl"]
# ///
# Documents produits à partir des gabarits Word/PowerPoint embarqués par python-docx/python-pptx
# (XML d'origine Microsoft : mc:Ignorable, w14/w15, etc.).
import base64, io, sys
from docx import Document
from docx.shared import Pt
from pptx import Presentation
from pptx.util import Inches

out = sys.argv[1]
png = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAIAAAACCAIAAAD91JpzAAAAFklEQVR4nGP8z8DAwMDAxMDAwMDAAAANHQEDasKb6QAAAABJRU5ErkJggg==")

d = Document()
d.core_properties.author = "Paulina Kowalski"
d.sections[0].header.paragraphs[0].text = "En-tête : Paulina Kowalski"
d.add_heading("Rapport", 1)
p = d.add_paragraph()
p.add_run("Rédigé par ").italic = True
p.add_run("Paulina ").bold = True
r = p.add_run("Kowalski"); r.bold = True; r.font.size = Pt(14)
p.add_run("  avec espaces  & entités <ok>.")
d.add_picture(io.BytesIO(png))
t = d.add_table(rows=2, cols=2)
for i, row in enumerate([["Nom", "E-mail"], ["Jean Dupont", "jean.dupont@example.fr"]]):
    for j, v in enumerate(row):
        t.cell(i, j).text = v
d.save(f"{out}/word.docx")

pr = Presentation()
s = pr.slides.add_slide(pr.slide_layouts[1])
s.shapes.title.text = "Présenté par Paulina Kowalski"
s.placeholders[1].text = "Contact : jean.dupont@example.fr"
s.shapes.add_picture(io.BytesIO(png), Inches(1), Inches(5))
s.notes_slide.notes_text_frame.text = "Appeler Jean Dupont"
pr.save(f"{out}/ppt.pptx")
